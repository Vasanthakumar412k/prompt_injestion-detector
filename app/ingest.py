"""
ingest.py – Document upload, text extraction, chunking, embedding, and storage.
Chunks are scanned at upload time; flagged chunks are quarantined before storage.
"""
from __future__ import annotations

import io
import logging
import re
import uuid
from pathlib import Path
from typing import List, Tuple

from sqlalchemy.orm import Session

from app.config import settings
from app.db import Chunk, Document, Permission, Role, User
from app.ensemble import run_ensemble
from app.db import ScanLog
from app.access_config import LEGACY_ROLE_MAP
from app.diagnostics import trace_step

logger = logging.getLogger(__name__)

# ── Lazy imports for heavy dependencies ─────────────────────────────────────

def _get_chroma():
    from app.vector_store import get_chroma_client

    return get_chroma_client()

def _get_collection():
    from app.vector_store import get_collection

    return get_collection()


# ── Text extraction ──────────────────────────────────────────────────────────


def extract_text(content: bytes, filename: str) -> str:
    """Extract plain text from txt, pdf, or docx bytes."""
    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        return content.decode("utf-8", errors="replace")

    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(io.BytesIO(content))
            return "\n".join(
                page.extract_text() or "" for page in reader.pages
            )
        except Exception as exc:
            logger.error("PDF extraction failed: %s", exc)
            return ""

    if ext == ".docx":
        try:
            import docx  # type: ignore

            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:
            logger.error("DOCX extraction failed: %s", exc)
            return ""

    raise ValueError(f"Unsupported file type: {ext}")


# ── Chunking ────────────────────────────────────────────────────────────────


def _rough_token_count(text: str) -> int:
    """Count model tokenizer tokens, including special tokens."""
    from app.embeddings import get_embedding_model

    model = get_embedding_model()
    return len(model.tokenizer.encode(text, add_special_tokens=True))


def _encode_content_tokens(tokenizer, text: str) -> List[int]:
    backend_tokenizer = getattr(tokenizer, "backend_tokenizer", None)
    if backend_tokenizer is not None:
        return backend_tokenizer.encode(text, add_special_tokens=False).ids
    return tokenizer.encode(text, add_special_tokens=False)


def chunk_text(text: str) -> List[str]:
    """
    Split text by embedding-model tokens without exceeding its input limit.
    """
    if not text.strip():
        return []

    from app.embeddings import get_embedding_model

    model = get_embedding_model()
    tokenizer = model.tokenizer
    max_seq_length = int(model.max_seq_length)
    special_tokens = tokenizer.num_special_tokens_to_add(pair=False)
    content_limit = max_seq_length - special_tokens
    size = min(settings.chunk_size_tokens, content_limit)
    overlap = min(settings.chunk_overlap_tokens, size - 1)
    if size < 1:
        raise ValueError("Embedding model max_seq_length cannot fit any content tokens.")

    token_ids = _encode_content_tokens(tokenizer, text)
    if not token_ids:
        return []

    chunks: List[str] = []
    start = 0
    while start < len(token_ids):
        end = min(start + size, len(token_ids))
        chunk_ids = token_ids[start:end]
        chunk = tokenizer.decode(chunk_ids, skip_special_tokens=True).strip()
        if len(tokenizer.encode(chunk, add_special_tokens=True)) > max_seq_length:
            raise ValueError(
                "Decoded embedding chunk exceeds model.max_seq_length; "
                "reduce chunk_size_tokens or use the configured model tokenizer."
            )
        chunks.append(chunk)
        if end == len(token_ids):
            break
        start += size - overlap

    return chunks


# ── Main ingest function ─────────────────────────────────────────────────────


def ingest_document(
    content: bytes,
    filename: str,
    content_type: str,
    uploader_id: int,
    allowed_role_names: List[str],
    db: Session,
    scan_on_ingest: bool = True,
    department_id: int | None = None,
    confidentiality: str = "internal",
) -> Tuple[Document, List[str]]:
    """
    Full ingest pipeline:
    1. Extract text.
    2. Chunk.
    3. Scan each chunk with ensemble (quarantine flagged chunks).
    4. Embed and store clean chunks in ChromaDB.
    5. Create SQL records for document, chunks, and permissions.

    Returns (Document, list_of_warning_messages).
    """
    warnings: List[str] = []
    trace_step("ingest", "start", uploader_id=uploader_id, byte_count=len(content), confidentiality=confidentiality)

    # ── Create document record ─────────────────────────────────────────────
    uploader = db.query(User).filter_by(id=uploader_id).first()
    doc = Document(
        filename=filename,
        content_type=content_type,
        uploader_id=uploader_id,
        owner_id=uploader_id,
        department_id=department_id if department_id is not None else (uploader.department_id if uploader else None),
        confidentiality=confidentiality,
    )
    db.add(doc)
    db.flush()  # get doc.id
    trace_step("ingest.document_record", "complete", document_id=doc.id, department_id=doc.department_id or 0)

    # ── Resolve roles and create permissions ───────────────────────────────
    effective_role_names = list(dict.fromkeys(
        LEGACY_ROLE_MAP.get(name, name) for name in allowed_role_names
    ))
    if confidentiality.lower() == "department" and "hr_manager" not in effective_role_names:
        effective_role_names.append("hr_manager")
    for role_name in effective_role_names:
        role = db.query(Role).filter(Role.name == role_name).first()
        if role:
            perm = Permission(
                role_id=role.id,
                document_id=doc.id,
                granted_by=uploader_id,
            )
            db.add(perm)
        else:
            warnings.append(f"Role '{role_name}' not found; permission skipped.")

    db.flush()

    # ── Extract and chunk ─────────────────────────────────────────────────
    try:
        trace_step("ingest.text_extraction", "start", document_id=doc.id)
        raw_text = extract_text(content, filename)
    except ValueError as exc:
        trace_step("ingest.text_extraction", "error", document_id=doc.id, error_type=type(exc).__name__)
        db.rollback()
        raise

    trace_step("ingest.text_extraction", "complete", document_id=doc.id, char_count=len(raw_text))
    chunks_text = chunk_text(raw_text)
    trace_step("ingest.chunking", "complete", document_id=doc.id, chunk_count=len(chunks_text))
    if not chunks_text:
        warnings.append("Document produced no text chunks.")
        db.commit()
        trace_step("ingest", "complete", document_id=doc.id, stored_chunks=0,
                   quarantined_chunks=0, warning_count=len(warnings))
        return doc, warnings

    # ── Scan + embed ──────────────────────────────────────────────────────
    trace_step("ingest.vector_store", "start", document_id=doc.id)
    try:
        collection = _get_collection()
    except Exception as exc:
        trace_step("ingest.vector_store", "error", document_id=doc.id, error_type=type(exc).__name__)
        db.rollback()
        raise
    allowed_roles_list = effective_role_names

    clean_ids: List[str] = []
    clean_texts: List[str] = []
    clean_metas: List[dict] = []

    for idx, chunk_text_str in enumerate(chunks_text):
        chroma_id = f"doc{doc.id}_chunk{idx}_{uuid.uuid4().hex[:8]}"

        # Security scan at upload time
        if scan_on_ingest:
            ensemble_result = run_ensemble(chunk_text_str, source="document")
            is_quarantined = ensemble_result.is_injection
            reason = ensemble_result.final_reason if is_quarantined else None
        else:
            is_quarantined = False
            reason = None
            trace_step("ingest.chunk_scan", "skipped", document_id=doc.id, chunk_index=idx)

        if scan_on_ingest:
            trace_step("ingest.chunk_scan", "complete", document_id=doc.id, chunk_index=idx,
                       verdict=ensemble_result.final_label, quarantined=is_quarantined)

        # Create SQL chunk record
        chunk_obj = Chunk(
            document_id=doc.id,
            chunk_index=idx,
            chroma_id=chroma_id,
            text_content=chunk_text_str,
            text_preview=chunk_text_str[:500],
            is_quarantined=is_quarantined,
            quarantine_reason=reason,
        )
        db.add(chunk_obj)
        db.flush()

        # Log the scan if performed
        if scan_on_ingest:
            _log_scan(db, chunk_text_str, ensemble_result, source="document", chunk_id=chunk_obj.id)

        if is_quarantined:
            warnings.append(
                f"Chunk {idx} quarantined at ingest: {reason}"
            )
            # Mark document for admin review if any chunk is poisoned
            doc.is_quarantined = True
            doc.quarantine_reason = (
                doc.quarantine_reason or ""
            ) + f"Chunk {idx} flagged: {reason}. "
            continue

        clean_ids.append(chroma_id)
        clean_texts.append(chunk_text_str)
        clean_metas.append({
            "doc_id": doc.id,
            "chunk_index": idx,
            "allowed_roles": ",".join(allowed_roles_list),
            "filename": filename,
            "department_id": doc.department_id or 0,
            "confidentiality": doc.confidentiality,
        })

    # ── Bulk upsert to ChromaDB ────────────────────────────────────────────
    if clean_ids:
        try:
            from app.embeddings import encode_texts

            collection.upsert(
                ids=clean_ids,
                documents=clean_texts,
                metadatas=clean_metas,
                embeddings=encode_texts(clean_texts),
            )
        except Exception as exc:
            trace_step("ingest.vector_store", "error", document_id=doc.id, error_type=type(exc).__name__)
            db.rollback()
            raise
        logger.info(
            "Ingested doc=%d: %d clean chunks, %d quarantined",
            doc.id, len(clean_ids), len(chunks_text) - len(clean_ids),
        )

    trace_step("ingest.vector_store", "complete", document_id=doc.id, stored_chunks=len(clean_ids))
    db.commit()
    trace_step("ingest", "complete", document_id=doc.id, stored_chunks=len(clean_ids),
               quarantined_chunks=len(chunks_text) - len(clean_ids), warning_count=len(warnings))
    return doc, warnings


def sync_department_role_metadata(db: Session) -> None:
    """Keep stored role metadata aligned with the department-guide read policy."""
    chunks = (
        db.query(Chunk)
        .join(Document, Chunk.document_id == Document.id)
        .filter(
            Document.confidentiality == "department",
            Chunk.status != "superseded",
        )
        .all()
    )
    if not chunks:
        return

    try:
        collection = _get_chroma().get_collection(name=settings.chroma_collection)
        records = collection.get(
            ids=[chunk.chroma_id for chunk in chunks],
            include=["metadatas"],
        )
    except Exception as exc:
        logger.error("Could not synchronize department-guide vector metadata: %s", exc)
        return

    metadata_by_id = dict(zip(records.get("ids", []), records.get("metadatas", [])))
    for chunk in chunks:
        metadata = metadata_by_id.get(chunk.chroma_id)
        if metadata is None:
            logger.warning(
                "No vector metadata found for department document chunk id=%d chroma_id=%s",
                chunk.id,
                chunk.chroma_id,
            )
            continue
        roles = {
            role.strip()
            for role in str(metadata.get("allowed_roles", "")).split(",")
            if role.strip()
        }
        if "hr_manager" in roles:
            continue
        roles.add("hr_manager")
        collection.update(
            ids=[chunk.chroma_id],
            metadatas=[{**metadata, "allowed_roles": ",".join(sorted(roles))}],
        )
        logger.info(
            "Updated department document role metadata document_id=%d chunk_id=%d",
            chunk.document_id,
            chunk.id,
        )


def _log_scan(db: Session, text: str, result, source: str, chunk_id: int | None = None, source_type: str = "live") -> None:
    """Persist an ensemble result to scan_logs."""
    log = ScanLog(
        source=source,
        source_type=source_type,
        text_preview=text[:500],
        deberta_label=result.deberta.label,
        deberta_score=result.deberta.score,
        deberta_latency_ms=result.deberta.latency_ms,
        deberta_error=result.deberta.error,
        promptguard_label=result.promptguard.label,
        promptguard_score=result.promptguard.score,
        promptguard_latency_ms=result.promptguard.latency_ms,
        promptguard_error=result.promptguard.error,
        gemini_used=result.gemini_used,
        gemini_label=result.gemini.label if result.gemini_used else None,
        gemini_reason=result.gemini.raw.get("reason") if result.gemini_used and isinstance(result.gemini.raw, dict) else None,
        gemini_confidence=result.gemini.score if result.gemini_used else None,
        gemini_error=result.gemini.error if result.gemini_used else None,
        final_label=result.final_label,
        final_reason=result.final_reason,
        chunk_id=chunk_id,
    )
    db.add(log)
